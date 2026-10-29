"""Formatted Excel report. Severity text is always present alongside color."""

from __future__ import annotations

from collections import defaultdict
from datetime import timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from scanner.models import SEVERITY_RANK, InstanceResult, Severity
from scanner.report_generator import ScanMetadata
from scanner.security_group_scanner import recommendation_for

TITLE = "Weekly AWS EC2 Security Exposure Report"
FINDING_NOTE = (
    "Findings represent potential security exposures based on AWS Security "
    "Group configuration. They do not prove exploitability."
)
NO_FINDINGS = "No Findings"

# Readable fills. Text labels are written in the same cells.
CRITICAL_FILL = "F8D0D8"
CRITICAL_FONT = "8B1E3F"
HIGH_FILL = "FDE6C8"
HIGH_FONT = "9A3412"
MEDIUM_FILL = "FBF3D0"
MEDIUM_FONT = "854D0E"
LOW_FILL = "DBEAFE"
LOW_FONT = "1D4ED8"
CLEAR_FILL = "DCFCE7"
CLEAR_FONT = "166534"
CLEAR_ROW_FILL = "F0FDF4"
HEADER_FILL = "F0F4F8"
TITLE_FONT = "102A43"

_SEVERITY_STYLE = {
    Severity.CRITICAL: (CRITICAL_FILL, CRITICAL_FONT, "FDECEE"),
    Severity.HIGH: (HIGH_FILL, HIGH_FONT, "FFF7ED"),
    Severity.MEDIUM: (MEDIUM_FILL, MEDIUM_FONT, "FEFCE8"),
    Severity.LOW: (LOW_FILL, LOW_FONT, "EFF6FF"),
}

_THIN = Border(
    left=Side(style="thin", color="BCCCDC"),
    right=Side(style="thin", color="BCCCDC"),
    top=Side(style="thin", color="BCCCDC"),
    bottom=Side(style="thin", color="BCCCDC"),
)
_BODY = Font(name="Calibri", size=11, color="1F2933")
_HEADER = Font(name="Calibri", size=11, bold=True, color=TITLE_FONT)
_WRAP = Alignment(vertical="center", wrap_text=True)
_LEFT = Alignment(vertical="center", horizontal="left", wrap_text=True)


def write_xlsx(
    path: Path, results: list[InstanceResult], metadata: ScanMetadata
) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    _write_summary(summary, results, metadata)
    _write_instance_summary(workbook.create_sheet("Instance Summary"), results)
    _write_details(workbook.create_sheet("Detailed Findings"), results)
    _write_recommendations(workbook.create_sheet("Recommendations"), results)
    workbook.save(path)


def _write_summary(
    sheet: Worksheet, results: list[InstanceResult], metadata: ScanMetadata
) -> None:
    when = metadata.generated_at.astimezone(timezone.utc)
    totals = {severity: 0 for severity in Severity}
    for result in results:
        for severity, count in result.severity_counts.items():
            totals[severity] += count
    affected = sum(1 for result in results if result.finding_count)
    sheet.merge_cells("A1:H1")
    sheet["A1"] = TITLE
    sheet["A1"].font = Font(name="Calibri", size=20, bold=True, color=TITLE_FONT)
    sheet.row_dimensions[1].height = 28
    if metadata.demo:
        sheet.merge_cells("A2:H2")
        sheet["A2"] = (
            "DEMO DATA. This report uses fictional instances and is not from AWS."
        )
        sheet["A2"].font = Font(name="Calibri", size=11, italic=True, color="9A3412")
    details = [
        ("Report Date", when.strftime("%Y-%m-%d")),
        ("Scan Time", when.strftime("%H:%M:%S UTC")),
        ("AWS Account", metadata.account_id),
        ("Region", metadata.region),
        ("Scanner Version", metadata.scanner_version),
        ("Scan Scope", metadata.scope),
    ]
    for offset, (label, value) in enumerate(details):
        row = 4 + offset
        _label_value(sheet, row, label, value)
        sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
    posture, posture_fill, posture_font = _posture(totals)
    sheet["A11"] = "Security Posture"
    sheet["A11"].font = Font(name="Calibri", size=12, bold=True, color=TITLE_FONT)
    sheet.merge_cells("A12:D12")
    posture_cell = sheet["A12"]
    posture_cell.value = posture
    posture_cell.font = Font(name="Calibri", size=14, bold=True, color=posture_font)
    posture_cell.fill = PatternFill(fill_type="solid", fgColor=posture_fill)
    posture_cell.alignment = _LEFT
    sheet.row_dimensions[12].height = 24
    metrics = [
        ("Total Running Instances", len(results), "F8FAFC", "1F2933"),
        ("Instances With Findings", affected, "FFF7ED", "9A3412"),
        ("Instances Without Findings", len(results) - affected, "F0FDF4", "166534"),
        ("Total Findings", sum(totals.values()), "F8FAFC", "1F2933"),
        ("Critical", totals[Severity.CRITICAL], CRITICAL_FILL, CRITICAL_FONT),
        ("High", totals[Severity.HIGH], HIGH_FILL, HIGH_FONT),
        ("Medium", totals[Severity.MEDIUM], MEDIUM_FILL, MEDIUM_FONT),
        ("Low", totals[Severity.LOW], LOW_FILL, LOW_FONT),
    ]
    for column, (label, value, fill, font_color) in enumerate(metrics, start=1):
        label_cell = sheet.cell(14, column, label)
        label_cell.font = Font(name="Calibri", size=9, bold=True, color="52606D")
        label_cell.alignment = Alignment(
            vertical="center", horizontal="center", wrap_text=True
        )
        value_cell = sheet.cell(15, column, value)
        value_cell.font = Font(name="Calibri", size=18, bold=True, color=font_color)
        value_cell.fill = PatternFill(fill_type="solid", fgColor=fill)
        value_cell.alignment = Alignment(vertical="center", horizontal="center")
        value_cell.border = _THIN
        label_cell.border = _THIN
        sheet.column_dimensions[get_column_letter(column)].width = 24
    sheet.row_dimensions[14].height = 30
    sheet.row_dimensions[15].height = 32
    sheet.merge_cells("A17:H17")
    note = sheet["A17"]
    note.value = FINDING_NOTE
    note.font = Font(name="Calibri", size=11, italic=True, color="334155")
    note.alignment = _WRAP
    sheet.row_dimensions[17].height = 32
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.page_setup.paperSize = sheet.PAPERSIZE_TABLOID
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_margins.left = 0.5
    sheet.page_margins.right = 0.5
    sheet.page_margins.top = 0.6
    sheet.page_margins.bottom = 0.6
    sheet.sheet_view.showGridLines = False
    sheet.oddHeader.left.text = TITLE
    sheet.print_options.horizontalCentered = True


def _write_instance_summary(sheet: Worksheet, results: list[InstanceResult]) -> None:
    headers = [
        "Status",
        "Instance",
        "Instance ID",
        "Team",
        "Private IP",
        "Public IP",
        "Findings",
        "Highest Severity",
    ]
    _write_headers(sheet, headers)
    for index, result in enumerate(results, start=2):
        instance = result.instance
        highest = result.highest_severity.value if result.highest_severity else "None"
        status = NO_FINDINGS if result.finding_count == 0 else "Findings"
        values = [
            status,
            instance.name,
            instance.instance_id,
            instance.team,
            instance.private_ip or "",
            instance.public_ip or "",
            result.finding_count,
            highest,
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(index, column, value)
            cell.font = _BODY
            cell.alignment = _LEFT
            cell.border = _THIN
        if result.finding_count == 0:
            _paint_row(sheet, index, len(headers), CLEAR_ROW_FILL)
            _paint_cell(sheet.cell(index, 1), CLEAR_FILL, CLEAR_FONT)
        else:
            strong, font_color, row_fill = _SEVERITY_STYLE[result.highest_severity]
            _paint_row(sheet, index, len(headers), row_fill)
            _paint_cell(sheet.cell(index, 1), strong, font_color)
            _paint_cell(sheet.cell(index, 8), strong, font_color)
        sheet.row_dimensions[index].height = 20
    _finish_table(sheet, headers, max(len(results) + 1, 1))


def _write_details(sheet: Worksheet, results: list[InstanceResult]) -> None:
    headers = [
        "Severity",
        "Instance",
        "Instance ID",
        "Team",
        "Port",
        "Protocol",
        "Source CIDR",
        "Security Group",
        "Security Group ID",
        "Vulnerability",
        "Description",
        "Recommendation",
    ]
    _write_headers(sheet, headers)
    row_index = 2
    for result in results:
        for finding in result.findings:
            vulnerability = finding.title
            if finding.rule_description:
                vulnerability = f"{vulnerability} ({finding.rule_description})"
            values = [
                finding.severity.value,
                finding.instance_name,
                finding.instance_id,
                result.instance.team,
                finding.port,
                finding.protocol,
                finding.source,
                finding.security_group_name,
                finding.security_group_id,
                vulnerability,
                finding.description,
                finding.recommendation,
            ]
            for column, value in enumerate(values, start=1):
                cell = sheet.cell(row_index, column, value)
                cell.font = _BODY
                cell.alignment = _LEFT
                cell.border = _THIN
            fill, font_color, _row = _SEVERITY_STYLE[finding.severity]
            _paint_cell(sheet.cell(row_index, 1), fill, font_color)
            sheet.row_dimensions[row_index].height = 48
            row_index += 1
    _finish_table(sheet, headers, max(row_index - 1, 1))


def _write_recommendations(sheet: Worksheet, results: list[InstanceResult]) -> None:
    headers = [
        "Issue",
        "Severity",
        "Affected Instances",
        "Finding Count",
        "Recommended Action",
    ]
    _write_headers(sheet, headers)
    grouped: dict[str, list[tuple[str, int, Severity]]] = defaultdict(list)
    for result in results:
        for finding in result.findings:
            grouped[_issue_name(finding.port)].append(
                (result.instance.instance_id, finding.port, finding.severity)
            )
    order = [
        "Public SSH exposure",
        "Public RDP exposure",
        "Public HTTP exposure",
        "Public HTTPS exposure",
        "Other monitored port exposure",
    ]
    row_index = 2
    for issue in order:
        items = grouped.get(issue, [])
        if not items:
            continue
        instances = sorted({instance_id for instance_id, _port, _severity in items})
        sample_port = items[0][1]
        severity = max(items, key=lambda item: SEVERITY_RANK[item[2]])[2]
        action = recommendation_for(sample_port, issue)
        values = [issue, severity.value, ", ".join(instances), len(items), action]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row_index, column, value)
            cell.font = _BODY
            cell.alignment = _LEFT
            cell.border = _THIN
        fill, font_color, _row = _SEVERITY_STYLE[severity]
        _paint_cell(sheet.cell(row_index, 2), fill, font_color)
        sheet.row_dimensions[row_index].height = 36
        row_index += 1
    if row_index == 2:
        sheet.cell(2, 1, "No potential security exposures were found.")
        sheet.cell(2, 1).font = _BODY
    _finish_table(sheet, headers, max(row_index - 1, 1))


def _posture(totals: dict[Severity, int]) -> tuple[str, str, str]:
    if totals[Severity.CRITICAL]:
        return "CRITICAL FINDINGS DETECTED", CRITICAL_FILL, CRITICAL_FONT
    if totals[Severity.HIGH]:
        return "HIGH OR MEDIUM FINDINGS DETECTED", HIGH_FILL, HIGH_FONT
    if totals[Severity.MEDIUM]:
        return "HIGH OR MEDIUM FINDINGS DETECTED", MEDIUM_FILL, MEDIUM_FONT
    if totals[Severity.LOW]:
        return "LOW FINDINGS DETECTED", LOW_FILL, LOW_FONT
    return "NO MONITORED PUBLIC EXPOSURES", CLEAR_FILL, CLEAR_FONT


def _issue_name(port: int) -> str:
    if port == 22:
        return "Public SSH exposure"
    if port == 3389:
        return "Public RDP exposure"
    if port == 80:
        return "Public HTTP exposure"
    if port == 443:
        return "Public HTTPS exposure"
    return "Other monitored port exposure"


def _label_value(sheet: Worksheet, row: int, label: str, value: object) -> None:
    label_cell = sheet.cell(row, 1, label)
    label_cell.font = _HEADER
    label_cell.alignment = _LEFT
    value_cell = sheet.cell(row, 2, value)
    value_cell.font = _BODY
    value_cell.alignment = _LEFT


def _write_headers(sheet: Worksheet, headers: list[str]) -> None:
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(1, column, header)
        cell.font = _HEADER
        cell.fill = PatternFill(fill_type="solid", fgColor=HEADER_FILL)
        cell.alignment = _LEFT
        cell.border = _THIN


def _finish_table(sheet: Worksheet, headers: list[str], last_row: int) -> None:
    last_column = get_column_letter(len(headers))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{last_column}{last_row}"
    widths = {
        "Instance": 22,
        "Instance ID": 22,
        "Team": 16,
        "Private IP": 16,
        "Public IP": 16,
        "Security Groups": 32,
        "Findings": 12,
        "Unique Ports": 14,
        "Highest Severity": 18,
        "Status": 16,
        "Severity": 14,
        "Port": 10,
        "Protocol": 12,
        "Source CIDR": 18,
        "Security Group": 22,
        "Security Group ID": 20,
        "Vulnerability": 36,
        "Description": 46,
        "Recommendation": 46,
        "Issue": 32,
        "Affected Instances": 36,
        "Finding Count": 16,
        "Recommended Action": 48,
    }
    for column, header in enumerate(headers, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = widths.get(
            header, 18
        )
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.paperSize = sheet.PAPERSIZE_TABLOID
    sheet.sheet_view.showGridLines = False


def _paint_row(sheet: Worksheet, row: int, columns: int, color: str) -> None:
    fill = PatternFill(fill_type="solid", fgColor=color)
    for column in range(1, columns + 1):
        sheet.cell(row, column).fill = fill


def _paint_cell(cell, fill_color: str, font_color: str) -> None:
    cell.fill = PatternFill(fill_type="solid", fgColor=fill_color)
    cell.font = Font(name="Calibri", size=11, bold=True, color=font_color)
