"""Excel workbook structure and severity colors. No AWS calls."""

from datetime import datetime, timezone

from factories import make_config, make_instance, permission, security_group
from openpyxl import load_workbook

from scanner.models import Severity
from scanner.report_generator import ScanMetadata, write_reports
from scanner.vulnerability_analyzer import analyze


def test_workbook_sheets_metrics_and_colors(tmp_path) -> None:
    exposed = make_instance(instance_id="i-exposed", name="web")
    clean = make_instance(
        instance_id="i-clean",
        name="quiet",
        public_ip=None,
        groups=(),
    )
    config = make_config(
        ports=(22, 80, 25),
        names={22: "SSH", 80: "HTTP", 25: "SMTP"},
        output_directory=str(tmp_path),
        formats=("xlsx",),
    )
    groups = {
        "sg-1": security_group(
            "sg-1",
            "web",
            [
                permission("tcp", 22, 22, ["0.0.0.0/0"]),
                permission("tcp", 80, 80, ["0.0.0.0/0"]),
                permission("tcp", 25, 25, ["0.0.0.0/0"]),
            ],
        )
    }
    results = analyze([exposed, clean], groups, config)
    written = write_reports(
        config,
        results,
        ScanMetadata(
            generated_at=datetime(2026, 10, 5, 18, 30, tzinfo=timezone.utc),
            region="ap-south-1",
            account_id="unavailable",
            scanner_version="0.1.0",
            scope="Running EC2 instances.",
            demo=False,
        ),
    )
    workbook = load_workbook(written[0])
    assert workbook.sheetnames == [
        "Summary",
        "Instance Summary",
        "Detailed Findings",
        "Recommendations",
    ]
    summary = workbook["Summary"]
    values = [
        cell
        for row in summary.iter_rows(max_row=20, max_col=8, values_only=True)
        for cell in row
        if cell
    ]
    assert "Weekly AWS EC2 Security Exposure Report" in values
    assert "Total Running Instances" in values
    assert "Instances With Findings" in values
    assert "Instances Without Findings" in values
    assert "Total Findings" in values
    assert "Critical" in values
    assert "High" in values
    assert "Medium" in values
    assert "Low" in values
    assert "CRITICAL FINDINGS DETECTED" in values
    assert any("do not prove exploitability" in str(value) for value in values)

    instances = workbook["Instance Summary"]
    assert instances.cell(1, 1).value == "Status"
    assert instances.cell(1, 8).value == "Highest Severity"
    rows = {instances.cell(row, 2).value: row for row in range(2, 4)}
    assert instances.cell(rows["quiet"], 1).value == "No Findings"
    assert instances.cell(rows["web"], 1).value == "Findings"
    assert instances.cell(rows["web"], 8).value == "CRITICAL"
    assert _color(instances.cell(rows["quiet"], 1)) == "DCFCE7"
    assert _color(instances.cell(rows["web"], 8)) == "F8D0D8"

    details = workbook["Detailed Findings"]
    assert details.cell(1, 1).value == "Severity"
    colors = {_color(details.cell(row, 1)) for row in range(2, 5)}
    assert "F8D0D8" in colors
    assert "FDE6C8" in colors
    assert "FBF3D0" in colors
    assert {details.cell(row, 1).value for row in range(2, 5)} == {
        Severity.CRITICAL.value,
        Severity.HIGH.value,
        Severity.MEDIUM.value,
    }

    recommendations = workbook["Recommendations"]
    assert recommendations.cell(1, 1).value == "Issue"
    issues = {recommendations.cell(row, 1).value for row in range(2, 6)}
    assert "Public SSH exposure" in issues
    assert "Public HTTP exposure" in issues
    assert "Other monitored port exposure" in issues


def test_details_headers_when_there_are_no_findings(tmp_path) -> None:
    written = write_reports(
        make_config(output_directory=str(tmp_path), formats=("xlsx",)),
        [],
        ScanMetadata(
            generated_at=datetime(2026, 10, 5, tzinfo=timezone.utc),
            region="ap-south-1",
            account_id="unavailable",
            scanner_version="0.1.0",
            scope="Running EC2 instances.",
            demo=False,
        ),
    )
    details = load_workbook(written[0])["Detailed Findings"]
    assert details.cell(1, 1).value == "Severity"
    assert details.cell(2, 1).value is None


def _color(cell) -> str:
    rgb = getattr(cell.fill.fgColor, "rgb", "") or ""
    return str(rgb)[-6:].upper()
